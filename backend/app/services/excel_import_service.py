"""Import des données depuis les fichiers Excel.

Colonnes attendues (première ligne = en-têtes, ordre libre, casse ignorée) :
- TypeEquipement.xlsx              : IDTypeEquipement, typeequipement
- NombreEquipement_TypeEquip.xlsx  : IDTypeEquipement, Nombre_Equipements, DateDebut, DateFin
  (la colonne IDNombreEquipement_TypeEquip est ignorée : l'id est attribué par la base)

Au démarrage, `importer_dossier` parcourt tous les .xlsx du dossier data/,
reconnaît leur type par les en-têtes et n'importe que les fichiers nouveaux
ou modifiés (empreinte SHA-256 mémorisée dans la table ImportFichier).

L'import est rejouable sans risque :
- types : ajoutés s'ils sont absents, renommés si le libellé a changé ;
- saisies : ignorées si identiques (même type + mêmes dates), rejetées si elles
  chevauchent une période existante (même règle que dans l'application),
  ajoutées sinon.
"""
import hashlib
import logging
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Equipement, ImportFichier, TypeEquipement
from app.schemas import EquipementIn
from app.services.equipement_service import find_chevauchement

logger = logging.getLogger(__name__)


class ExcelFormatError(Exception):
    pass


@dataclass
class RapportImport:
    types_ajoutes: int = 0
    types_renommes: int = 0
    saisies_ajoutees: int = 0
    saisies_deja_presentes: int = 0
    rejets: list[str] = field(default_factory=list)
    types_lus: bool = False
    saisies_lues: bool = False

    def resume(self) -> str:
        lignes = []
        if self.types_lus:
            lignes.append(f"Types : {self.types_ajoutes} ajouté(s), {self.types_renommes} renommé(s)")
        if self.saisies_lues:
            lignes.append(
                f"Saisies : {self.saisies_ajoutees} ajoutée(s), {self.saisies_deja_presentes} déjà présente(s), "
                f"{len(self.rejets)} rejetée(s)"
            )
        lignes += [f"  - {r}" for r in self.rejets]
        return "\n".join(lignes)


def _lire_lignes(path: str | Path, colonnes: list[str]) -> list[tuple[int, dict]]:
    """Lit la première feuille et renvoie [(numéro de ligne Excel, {colonne: valeur})]."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Fichier introuvable : {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        entetes = [str(h).strip().lower() if h is not None else "" for h in next(rows, [])]
        manquantes = [c for c in colonnes if c.lower() not in entetes]
        if manquantes:
            raise ExcelFormatError(f"{path.name} : colonne(s) manquante(s) : {', '.join(manquantes)}")
        index = {c: entetes.index(c.lower()) for c in colonnes}
        resultat = []
        for num, row in enumerate(rows, start=2):
            if row is None or all(v is None for v in row):
                continue  # ligne vide
            resultat.append((num, {c: row[i] if i < len(row) else None for c, i in index.items()}))
        return resultat
    finally:
        wb.close()


def _en_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str) and v.strip():
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def importer_types(db: Session, path: str | Path, rapport: RapportImport) -> None:
    rapport.types_lus = True
    for num, r in _lire_lignes(path, ["IDTypeEquipement", "typeequipement"]):
        try:
            type_id = int(r["IDTypeEquipement"])
        except (TypeError, ValueError):
            rapport.rejets.append(f"{Path(path).name} ligne {num} : IDTypeEquipement invalide ({r['IDTypeEquipement']!r})")
            continue
        nom = str(r["typeequipement"]).strip() if r["typeequipement"] is not None else None
        existant = db.get(TypeEquipement, type_id)
        if existant is None:
            db.add(TypeEquipement(id=type_id, nom=nom))
            rapport.types_ajoutes += 1
        elif existant.nom != nom:
            existant.nom = nom
            rapport.types_renommes += 1
    db.flush()


def importer_nombres(db: Session, path: str | Path, rapport: RapportImport) -> None:
    rapport.saisies_lues = True
    nom_fichier = Path(path).name
    colonnes = ["IDTypeEquipement", "Nombre_Equipements", "DateDebut", "DateFin"]
    for num, r in _lire_lignes(path, colonnes):
        ref = f"{nom_fichier} ligne {num}"
        try:
            data = EquipementIn(
                type_equipement_id=int(r["IDTypeEquipement"]),
                nombre=int(r["Nombre_Equipements"]),
                date_debut=_en_date(r["DateDebut"]),
                date_fin=_en_date(r["DateFin"]),
            )
        except (TypeError, ValueError, ValidationError) as e:
            detail = e.errors()[0]["msg"].removeprefix("Value error, ") if isinstance(e, ValidationError) else "valeur invalide"
            rapport.rejets.append(f"{ref} : {detail}")
            continue

        if db.get(TypeEquipement, data.type_equipement_id) is None:
            rapport.rejets.append(f"{ref} : type d'équipement {data.type_equipement_id} inconnu")
            continue

        identique = db.scalars(
            select(Equipement).where(
                Equipement.type_equipement_id == data.type_equipement_id,
                Equipement.date_debut == data.date_debut,
                Equipement.date_fin == data.date_fin,
            )
        ).first()
        if identique:
            rapport.saisies_deja_presentes += 1
            continue

        conflit = find_chevauchement(db, data)
        if conflit:
            rapport.rejets.append(
                f"{ref} : chevauche la période du {conflit.date_debut:%d/%m/%Y} au {conflit.date_fin:%d/%m/%Y} "
                f"(type {data.type_equipement_id})"
            )
            continue

        db.add(Equipement(**data.model_dump()))
        db.flush()  # pour que les lignes suivantes du fichier voient celle-ci
        rapport.saisies_ajoutees += 1


def importer_excel(db: Session, types_path: str | Path | None, nombres_path: str | Path | None) -> RapportImport:
    """Importe les types puis les saisies, en une seule transaction."""
    rapport = RapportImport()
    try:
        if types_path:
            importer_types(db, types_path, rapport)
        if nombres_path:
            importer_nombres(db, nombres_path, rapport)
        db.commit()
    except Exception:
        db.rollback()
        raise
    return rapport


# --------------------------------------------------------------------------
# Import automatique d'un dossier
# --------------------------------------------------------------------------

COLONNES_TYPES = {"idtypeequipement", "typeequipement"}
COLONNES_NOMBRES = {"idtypeequipement", "nombre_equipements", "datedebut", "datefin"}


def detecter_type_fichier(path: Path) -> str | None:
    """'types', 'nombres' ou None, d'après les en-têtes de la première feuille."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        premiere = next(wb.worksheets[0].iter_rows(values_only=True, max_row=1), ())
    finally:
        wb.close()
    entetes = {str(h).strip().lower() for h in premiere if h is not None}
    if COLONNES_NOMBRES <= entetes:
        return "nombres"
    if COLONNES_TYPES <= entetes:
        return "types"
    return None


def _empreinte(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _memoriser(db: Session, nom: str, empreinte: str, resume: str) -> None:
    ligne = db.scalars(select(ImportFichier).where(ImportFichier.nom_fichier == nom)).first()
    if ligne is None:
        db.add(ImportFichier(nom_fichier=nom, empreinte=empreinte, resume=resume))
    else:
        ligne.empreinte, ligne.resume, ligne.date_import = empreinte, resume, datetime.utcnow()
    db.commit()


def importer_dossier(db: Session, dossier: str | Path, forcer: bool = False) -> list[tuple[str, str]]:
    """Importe les fichiers Excel nouveaux ou modifiés du dossier.

    Les fichiers de types sont traités avant les fichiers de saisies.
    Chaque fichier est importé dans sa propre transaction : un fichier en
    erreur n'empêche pas les autres et sera retenté au prochain démarrage.
    Renvoie [(nom du fichier, résumé)] pour les fichiers traités.
    """
    dossier = Path(dossier)
    if not dossier.is_dir():
        logger.warning("Dossier d'import introuvable : %s", dossier)
        return []

    deja = {f.nom_fichier: f.empreinte for f in db.scalars(select(ImportFichier))}
    a_traiter: dict[str, list[tuple[Path, str]]] = {"types": [], "nombres": []}
    resultats: list[tuple[str, str]] = []

    # "~$..." : fichiers de verrouillage créés par Excel quand un classeur est ouvert
    for path in sorted(p for p in dossier.glob("*.xlsx") if not p.name.startswith("~$")):
        empreinte = _empreinte(path)
        if not forcer and deja.get(path.name) == empreinte:
            continue  # inchangé depuis le dernier import
        try:
            genre = detecter_type_fichier(path)
        except Exception as e:  # fichier corrompu ou illisible
            resultats.append((path.name, f"Illisible : {e}"))
            continue
        if genre is None:
            resume = ("Format non reconnu, fichier ignoré. En-têtes attendus : "
                      "IDTypeEquipement + typeequipement (types) ou "
                      "IDTypeEquipement + Nombre_Equipements + DateDebut + DateFin (saisies)")
            _memoriser(db, path.name, empreinte, resume)
            resultats.append((path.name, resume))
            continue
        a_traiter[genre].append((path, empreinte))

    for genre in ("types", "nombres"):
        for path, empreinte in a_traiter[genre]:
            try:
                rapport = importer_excel(
                    db, path if genre == "types" else None, path if genre == "nombres" else None
                )
            except Exception as e:
                db.rollback()
                resultats.append((path.name, f"Erreur, sera retenté au prochain démarrage : {e}"))
                continue
            resume = rapport.resume()
            _memoriser(db, path.name, empreinte, resume)
            resultats.append((path.name, resume))

    return resultats
